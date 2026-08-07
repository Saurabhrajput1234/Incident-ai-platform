"""
Excel/CSV parser for the Shift Roster module.

Expected Excel sheet: "Roster"
Columns:
  0: Assignment Group
  1: Assigned To
  2: Email
  3: Shift (default shift)
  4: <month header>  e.g. "2026-07-01 to 2026-07-31"
  5: Level
  6..36: Day 1 to Day 31 shift values

Output:
  ParsedRoster dataclass with:
    - roster_start_date / roster_end_date
    - rows: list of ParsedEngineerRow
"""
import io
import csv
from datetime import date, datetime
from dataclasses import dataclass, field


@dataclass
class ParsedEngineerRow:
    assignment_group: str
    assigned_to: str
    email: str
    default_shift: str | None
    level: str | None
    # date → shift_code mapping (only days with a value are included)
    daily_shifts: dict[date, str] = field(default_factory=dict)


@dataclass
class ParsedRoster:
    roster_start_date: date
    roster_end_date: date
    rows: list[ParsedEngineerRow] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)


# Column indices (0-based)
COL_GROUP = 0
COL_NAME = 1
COL_EMAIL = 2
COL_DEFAULT_SHIFT = 3
COL_MONTH_HEADER = 4
COL_LEVEL = 5
COL_DAYS_START = 6


def _parse_month_header(value: str) -> tuple[date, date]:
    """
    Parse '2026-07-01 to 2026-07-31' into (start_date, end_date).
    Raises ValueError on invalid format.
    """
    try:
        parts = str(value).split(" to ")
        start = datetime.strptime(parts[0].strip(), "%Y-%m-%d").date()
        end = datetime.strptime(parts[1].strip(), "%Y-%m-%d").date()
        return start, end
    except Exception:
        raise ValueError(f"Cannot parse month header: '{value}'. Expected 'YYYY-MM-DD to YYYY-MM-DD'")


def parse_excel(file_bytes: bytes) -> ParsedRoster:
    """Parse .xlsx file into ParsedRoster."""
    from openpyxl import load_workbook

    wb = load_workbook(filename=io.BytesIO(file_bytes), read_only=True, data_only=True)

    if "Roster" not in wb.sheetnames:
        raise ValueError("Excel file must contain a sheet named 'Roster'")

    ws = wb["Roster"]
    rows = list(ws.iter_rows(values_only=True))

    if not rows:
        raise ValueError("Roster sheet is empty")

    # Parse header row
    header = rows[0]
    try:
        start_date, end_date = _parse_month_header(header[COL_MONTH_HEADER])
    except ValueError as e:
        raise ValueError(str(e))

    result = ParsedRoster(roster_start_date=start_date, roster_end_date=end_date)

    for row_idx, row in enumerate(rows[1:], start=2):
        if not any(row):
            continue  # skip blank rows

        assigned_to = row[COL_NAME]
        if not assigned_to:
            result.errors.append(f"Row {row_idx}: missing 'Assigned To' — skipped")
            continue

        email = row[COL_EMAIL]
        if not email:
            result.errors.append(f"Row {row_idx}: missing 'Email' for '{assigned_to}' — skipped")
            continue

        # Build daily shifts dict: {date: shift_code}
        daily_shifts: dict[date, str] = {}
        for day_num in range(1, 32):
            col_idx = COL_DAYS_START + (day_num - 1)
            val = row[col_idx] if col_idx < len(row) else None
            if val:
                try:
                    roster_date = date(start_date.year, start_date.month, day_num)
                    daily_shifts[roster_date] = str(val).strip()
                except ValueError:
                    pass  # day doesn't exist in this month (e.g. Feb 30)

        result.rows.append(ParsedEngineerRow(
            assignment_group=str(row[COL_GROUP]).strip() if row[COL_GROUP] else "",
            assigned_to=str(assigned_to).strip(),
            email=str(email).strip().lower(),
            default_shift=str(row[COL_DEFAULT_SHIFT]).strip() if row[COL_DEFAULT_SHIFT] else None,
            level=str(row[COL_LEVEL]).strip() if row[COL_LEVEL] else None,
            daily_shifts=daily_shifts,
        ))

    return result


def parse_csv(file_bytes: bytes) -> ParsedRoster:
    """
    Parse .csv file. Expected columns (same as Excel):
    Assignment Group, Assigned To, Email, Shift, Month, Level, 1, 2, ..., 31
    """
    content = file_bytes.decode("utf-8-sig")
    reader = csv.reader(io.StringIO(content))
    rows = list(reader)

    if not rows:
        raise ValueError("CSV file is empty")

    header = rows[0]
    try:
        start_date, end_date = _parse_month_header(header[COL_MONTH_HEADER])
    except (ValueError, IndexError) as e:
        raise ValueError(str(e))

    result = ParsedRoster(roster_start_date=start_date, roster_end_date=end_date)

    for row_idx, row in enumerate(rows[1:], start=2):
        if not any(row):
            continue

        assigned_to = row[COL_NAME] if len(row) > COL_NAME else None
        if not assigned_to:
            result.errors.append(f"Row {row_idx}: missing 'Assigned To' — skipped")
            continue

        email = row[COL_EMAIL] if len(row) > COL_EMAIL else None
        if not email:
            result.errors.append(f"Row {row_idx}: missing 'Email' for '{assigned_to}' — skipped")
            continue

        daily_shifts: dict[date, str] = {}
        for day_num in range(1, 32):
            col_idx = COL_DAYS_START + (day_num - 1)
            val = row[col_idx] if col_idx < len(row) else None
            if val and val.strip():
                try:
                    roster_date = date(start_date.year, start_date.month, day_num)
                    daily_shifts[roster_date] = val.strip()
                except ValueError:
                    pass

        result.rows.append(ParsedEngineerRow(
            assignment_group=row[COL_GROUP].strip() if len(row) > COL_GROUP else "",
            assigned_to=assigned_to.strip(),
            email=email.strip().lower(),
            default_shift=row[COL_DEFAULT_SHIFT].strip() if len(row) > COL_DEFAULT_SHIFT and row[COL_DEFAULT_SHIFT] else None,
            level=row[COL_LEVEL].strip() if len(row) > COL_LEVEL and row[COL_LEVEL] else None,
            daily_shifts=daily_shifts,
        ))

    return result
